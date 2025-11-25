from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
# Use Title and Content layout (usually layout 1)
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)

# Title
title = slide.shapes.title
title.text = "Databaseoversikt — nøkkelpunkter"

# Body (bullets)
body = slide.shapes.placeholders[1]
tf = body.text_frame
tf.word_wrap = True

bullets = [
    "ArtService: Arts — kunstobjekter (Id, Title, Price, SellerId)",
    "AuctionService: Auctions & Bids — auksjoner og bud",
    "UserService: Profiles — bruker‑metadata (UserId, KeycloakId, Email)",
    "PaymentService: Transactions — betalinger og leverandør‑referanser (Amount, Status)",
    "AdminService: AdminLogs — audit / admin‑hendelser",
    "Viktig: Egen DB per tjeneste; GUID‑felt er logiske pekere; Keycloak for auth"
]

# Add bullets
for i, b in enumerate(bullets):
    if i == 0:
        p = tf.paragraphs[0]
        p.text = b
    else:
        p = tf.add_paragraph()
        p.text = b
    p.level = 0
    # styling
    for run in p.runs:
        run.font.size = Pt(18)

# Footer: add a textbox at bottom
left = Inches(0.3)
width = Inches(9)
height = Inches(0.5)
top = Inches(6.6)
textbox = slide.shapes.add_textbox(left, top, width, height)
tf2 = textbox.text_frame
p = tf2.paragraphs[0]
p.text = "Mikrotjenester m/ egne DB-er · GUID‑referanser · Keycloak auth"
p.font.size = Pt(12)
p.font.italic = True
p.font.color.rgb = RGBColor(100, 100, 100)

# Speaker notes
notes_slide = slide.notes_slide
notes = notes_slide.notes_text_frame
notes.text = "Tjeneste-eide databaser med logiske GUID‑referanser; bruk hendelsesmønstre (sagas/events) for sikre tverrgående transaksjoner." 

# Save
prs.save('docs/DatabaseSlide.pptx')
print('Saved docs/DatabaseSlide.pptx')

